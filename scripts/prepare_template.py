#!/usr/bin/env python3
"""生成精简模板：仅保留四大题型说明、题目区、答案区。"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from doc_styles import set_run_font

OUT = ROOT / "templates" / "template.docx"

SECTION_HEADERS = {
    "single": "一、单选题，以下各题有多个选项，其中只有一个选项是正确的，请选择正确答案(本大题满分30分,每小题1分)",
    "multiple": "二、多选题，以下各题有多个选项，其中有二个或多个选项是正确的，请选择正确答案(本大题满分20分,每小题1分)",
    "judge": "三、判断题，以下各题只有对错两个选项(本大题满分20分,每小题1分)",
    "short": "四、简答题(本大题满分30分,每小题6分)",
}

QUESTION_PLACEHOLDERS = {
    "single": "{{p single_questions }}",
    "multiple": "{{p multiple_questions }}",
    "judge": "{{p judge_questions }}",
    "short": "{{p short_questions }}",
}

ANSWER_SECTIONS = [
    ("一、单选题", "{{p single_answer_lines }}"),
    ("二、多选题", "{{p multiple_answer_lines }}"),
    ("三、判断题", "{{p judge_answer_lines }}"),
    ("四、简答题", "{{p short_answer_lines }}"),
]


def _add_section_header(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, bold=True)


def _add_placeholder(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def build_minimal_template() -> None:
    doc = Document()

    for key in ("single", "multiple", "judge", "short"):
        _add_section_header(doc, SECTION_HEADERS[key])
        _add_placeholder(doc, QUESTION_PLACEHOLDERS[key])

    answer_title = doc.add_paragraph()
    answer_title.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    answer_title.add_run().add_break(WD_BREAK.PAGE)
    title_run = answer_title.add_run("{% if paper_title %}《{{ paper_title }}》{% endif %}答案")
    set_run_font(title_run, bold=True)

    for section_title, placeholder in ANSWER_SECTIONS:
        _add_section_header(doc, section_title)
        _add_placeholder(doc, placeholder)

    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for para in section.header.paragraphs:
            para.text = ""
        for para in section.footer.paragraphs:
            para.text = ""

    doc.save(str(OUT))
    print(f"精简模板已生成: {OUT}，段落数 {len(doc.paragraphs)}")


def prepare_template() -> None:
    build_minimal_template()


if __name__ == "__main__":
    prepare_template()
