"""修复样板转 docx 后的版式问题。"""

from __future__ import annotations

import re

from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

PAGE_PATTERN = re.compile(r"=page|SectionPages|第\s*页（共", re.IGNORECASE)
JUNK_PATTERN = re.compile(
    r"(=page|SectionPages|第\s*\d*\s*页（共|"
    r"^-+\s*装\s*-+|"
    r"^○|"
    r"^(单\s*位|姓\s*名|考\s*号|座位号|部职别)$)"
)


def _delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _paragraph_has_page_break(paragraph: Paragraph) -> bool:
    xml = paragraph._element.xml
    return 'w:type="page"' in xml or "lastRenderedPageBreak" in xml


def ensure_page_break_before(paragraph: Paragraph) -> None:
    if _paragraph_has_page_break(paragraph):
        return
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)
    br_run = paragraph.runs[-1]._element
    paragraph._element.remove(br_run)
    paragraph._element.insert(0, br_run)


def set_paragraph_text_keep_break(paragraph: Paragraph, text: str) -> None:
    has_break = _paragraph_has_page_break(paragraph)
    element = paragraph._element
    for child in list(element):
        if child.tag.endswith("}pPr"):
            continue
        element.remove(child)
    if has_break:
        br_run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        br_run.append(br)
        element.append(br_run)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    element.append(run)


def is_trailing_junk(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    return bool(JUNK_PATTERN.search(stripped))


def trim_trailing_appendix(doc) -> int:
    """删除答案区之后误转入正文的页码段、装订线、信息表等。"""
    answer_idx = None
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith("《") and paragraph.text.strip().endswith("答案"):
            answer_idx = idx
            break
    if answer_idx is None:
        return 0

    removed = 0
    while len(doc.paragraphs) - 1 > answer_idx:
        idx = len(doc.paragraphs) - 1
        paragraph = doc.paragraphs[idx]
        text = paragraph.text.strip()
        if is_trailing_junk(text):
            _delete_paragraph(paragraph)
            removed += 1
            continue
        if not text:
            _delete_paragraph(paragraph)
            removed += 1
            continue
        break
    return removed
