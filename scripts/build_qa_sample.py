#!/usr/bin/env python3
"""从样板卷提取简答题，生成标准格式 sample_data/问答.xls。"""

from pathlib import Path

import xlwt
from docx import Document

import re

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "templates" / "template_raw.docx"
OUT = ROOT / "sample_data" / "问答.xls"
ANSWER_ITEM = re.compile(r"^\d+\.")


def _collect_stems(doc: Document) -> list[str]:
    stems: list[str] = []
    in_section = False
    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("四、简答题") and "本大题" in text:
            in_section = True
            continue
        if not in_section:
            continue
        if text.startswith("《"):
            break
        if text and text[0].isdigit() and ". " in text:
            stems.append(text.split(". ", 1)[1])
    return stems


def _collect_answers(doc: Document) -> list[str]:
    answers: list[str] = []
    in_answer_block = False
    in_section = False
    current: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("《") and "答案" in text:
            in_answer_block = True
            continue
        if not in_answer_block:
            continue
        if text == "四、简答题":
            in_section = True
            continue
        if not in_section:
            continue
        if text.startswith("《"):
            if current:
                answers.append("\n".join(current))
            break
        if text and ANSWER_ITEM.match(text):
            if current:
                answers.append("\n".join(current))
            current = [text.split(".", 1)[1].strip()]
            continue
        if text and current:
            current.append(text)
    if current:
        answers.append("\n".join(current))
    return answers


def main() -> None:
    doc = Document(str(RAW))
    stems = _collect_stems(doc)
    answers = _collect_answers(doc)
    count = min(len(stems), len(answers))
    if count < 5:
        raise SystemExit(f"简答题不足 5 道: stems={len(stems)} answers={len(answers)}")

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet0")
    sheet.write(0, 0, "问答类")
    sheet.write(0, 1, "710型船柴油机专业操作技能高级问答")
    headers = ["编号", "题目", "正确答案"] + [f"备选答案{i}" for i in range(1, 7)]
    for col, header in enumerate(headers):
        sheet.write(1, col, header)

    for row in range(count):
        sheet.write(row + 2, 0, float(row + 1))
        sheet.write(row + 2, 1, stems[row])
        sheet.write(row + 2, 2, answers[row])

    workbook.save(str(OUT))
    print(f"已生成 {OUT}，共 {count} 道简答题")


if __name__ == "__main__":
    main()
