"""将 textutil 转换后的单栏模板重建为图1版式（双栏+得分表+装订线）。"""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docx.text.paragraph import Paragraph


def _delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _set_run_font(run, name: str = "宋体", size_half: int = 21, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size_half / 2)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), "Times New Roman")
    for tag in ("sz", "szCs"):
        for old in r_pr.findall(qn(f"w:{tag}")):
            r_pr.remove(old)
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:val"), str(size_half))
        r_pr.append(node)


def _find_paragraph(doc: Document, predicate, start: int = 0) -> int:
    for idx in range(start, len(doc.paragraphs)):
        if predicate(doc.paragraphs[idx].text.strip()):
            return idx
    raise ValueError("未找到目标段落")


def _insert_section_break(paragraph: Paragraph, num_cols: int, page_break: bool = False) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    for old in p_pr.findall(qn("w:sectPr")):
        p_pr.remove(old)
    sect_pr = OxmlElement("w:sectPr")
    if page_break:
        break_type = OxmlElement("w:type")
        break_type.set(qn("w:val"), "nextPage")
        sect_pr.append(break_type)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num_cols))
    cols.set(qn("w:space"), "720")
    cols.set(qn("w:equalWidth"), "1")
    sect_pr.append(cols)
    pg_mar = OxmlElement("w:pgMar")
    pg_mar.set(qn("w:top"), "1440")
    pg_mar.set(qn("w:right"), "1080")
    pg_mar.set(qn("w:bottom"), "1440")
    pg_mar.set(qn("w:left"), "1800")
    pg_mar.set(qn("w:header"), "720")
    pg_mar.set(qn("w:footer"), "720")
    pg_mar.set(qn("w:gutter"), "0")
    sect_pr.append(pg_mar)
    p_pr.append(sect_pr)


def _style_table_cell(cell, text: str, bold: bool = False, center: bool = True) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER if center else WD_PARAGRAPH_ALIGNMENT.LEFT
    run = para.add_run(text)
    _set_run_font(run, "宋体", 21, bold=bold)


def _replace_score_list_with_table(doc: Document) -> None:
    """将竖排「题型/单选题/…」替换为 7 列表格。"""
    start = _find_paragraph(doc, lambda t: t == "题型")
    end = _find_paragraph(doc, lambda t: t == "得分", start=start) + 1
    anchor = doc.paragraphs[start - 1] if start > 0 else doc.paragraphs[0]

    for idx in range(end - 1, start - 1, -1):
        _delete_paragraph(doc.paragraphs[idx])

    table = doc.add_table(rows=2, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["题型", "单选题", "多选题", "判断题", "简答题", "总分", "得分"]
    for col, header in enumerate(headers):
        _style_table_cell(table.rows[0].cells[col], header, bold=True)
    _style_table_cell(table.rows[1].cells[0], "得分", bold=False)
    for col in range(1, 7):
        _style_table_cell(table.rows[1].cells[col], "", bold=False)

    anchor._element.addnext(table._tbl)
    tbl = anchor._element.getnext()
    empty_p = OxmlElement("w:p")
    tbl.addnext(empty_p)


def _add_seat_number_table(doc: Document) -> None:
    title_idx = 0
    title_p = doc.paragraphs[title_idx]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    _style_table_cell(table.rows[0].cells[0], "座位号", bold=False)
    _style_table_cell(table.rows[0].cells[1], "", bold=False)
    title_p._element.addprevious(table._tbl)


def _setup_binding_header(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        for p in list(header.paragraphs):
            p.text = ""
        for tbl in list(header.tables):
            tbl._element.getparent().remove(tbl._element)
        table = header.add_table(rows=1, cols=2, width=Inches(6.5))
        table.columns[0].width = Cm(1.8)
        left = table.rows[0].cells[0]
        left.text = ""
        lines = ["单  位", "", "姓  名", "", "考  号", "", "装", "订", "线"]
        for i, line in enumerate(lines):
            para = left.paragraphs[0] if i == 0 else left.add_paragraph()
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = para.add_run(line)
            _set_run_font(run, "宋体", 18)
        right = table.rows[0].cells[1]
        right.text = ""


def _apply_column_sections(doc: Document) -> None:
    single_header = _find_paragraph(
        doc,
        lambda t: t.startswith("一、单选题，以下各题有多个选项，其中只有一个选项是正确的"),
    )
    answer_title = _find_paragraph(doc, lambda t: t.startswith("《") and t.endswith("答案"))

    if single_header > 0:
        _insert_section_break(doc.paragraphs[single_header - 1], num_cols=2, page_break=False)
    _insert_section_break(doc.paragraphs[answer_title - 1], num_cols=1, page_break=True)


def apply_exam_layout(doc: Document) -> None:
    """应用图1版式：得分表、座位号、双栏正文、页眉装订区。"""
    _replace_score_list_with_table(doc)
    _add_seat_number_table(doc)
    _apply_column_sections(doc)
    _setup_binding_header(doc)
