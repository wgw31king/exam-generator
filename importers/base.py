"""Excel 读取：兼容标准 .xls、.xlsx、WPS/麒麟表格，必要时自动转换。"""

from __future__ import annotations

import hashlib
import html.parser
import shutil
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import xlrd
import xlwt
from xlrd.biffh import XLRDError


class SheetAdapter:
    """统一 sheet 接口，供 importers 使用。"""

    def __init__(self, rows: list[list[Any]]):
        self._rows = rows

    @property
    def nrows(self) -> int:
        return len(self._rows)

    @property
    def ncols(self) -> int:
        return max((len(row) for row in self._rows), default=0)

    def cell_value(self, row: int, col: int) -> Any:
        if row >= self.nrows:
            return ""
        row_data = self._rows[row]
        if col >= len(row_data):
            return ""
        return row_data[col]


def _app_dir() -> Path:
    return Path(__file__).resolve().parent.parent


def _read_header(path: Path, size: int = 8) -> bytes:
    with path.open("rb") as handle:
        return handle.read(size)


def is_standard_xls(path: Path) -> bool:
    return _read_header(path).startswith(b"\xd0\xcf\x11\xe0")


def is_zip_spreadsheet(path: Path) -> bool:
    return _read_header(path).startswith(b"PK")


def is_wps_binary(path: Path) -> bool:
    header = _read_header(path, 2)
    return header in (b"\xc1\x12", b"\xd0\x12")


def _cache_dir() -> Path:
    directory = _app_dir() / ".cache" / "converted"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _cache_path(source: Path) -> Path:
    stat = source.stat()
    digest = hashlib.sha1(
        f"{source.resolve()}:{stat.st_mtime_ns}:{stat.st_size}".encode()
    ).hexdigest()[:16]
    return _cache_dir() / f"{digest}_{source.stem}.xls"


def _normalize_rows(rows: list[list[Any]]) -> list[list[Any]]:
    if not rows:
        return [[]]
    width = max(len(row) for row in rows)
    return [list(row) + [""] * (width - len(row)) for row in rows]


def _write_standard_xls(rows: list[list[Any]], destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Sheet1")
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            if value is None or value == "":
                continue
            if isinstance(value, float) and value == int(value):
                sheet.write(row_idx, col_idx, int(value))
            else:
                sheet.write(row_idx, col_idx, str(value))
    workbook.save(str(destination))
    return destination


def _rows_to_sheet(rows: list[list[Any]]) -> SheetAdapter:
    return SheetAdapter(_normalize_rows(rows))


def _load_calamine(path: Path) -> SheetAdapter | None:
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None
    try:
        workbook = CalamineWorkbook.from_path(str(path))
        rows = workbook.get_sheet_by_index(0).to_python(skip_empty_area=False)
        return _rows_to_sheet(rows)
    except Exception:
        return None


def _load_openpyxl(path: Path) -> SheetAdapter:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_to_sheet(rows)


def _try_openpyxl_variants(path: Path) -> SheetAdapter | None:
    attempts = [path]
    if is_zip_spreadsheet(path):
        attempts.extend(
            [
                path.with_suffix(".xlsx"),
                path.with_suffix(".et"),
            ]
        )
    for candidate in attempts:
        if not candidate.exists():
            continue
        try:
            return _load_openpyxl(candidate)
        except Exception:
            continue

    if is_zip_spreadsheet(path):
        return None

    with tempfile.TemporaryDirectory() as tmp:
        for suffix in (".xlsx", ".et", ".xlsm"):
            copied = Path(tmp) / f"{path.stem}{suffix}"
            shutil.copy2(path, copied)
            try:
                return _load_openpyxl(copied)
            except Exception:
                continue
    return None


def _find_libreoffice() -> list[str] | None:
    candidates = [
        "libreoffice",
        "soffice",
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        "/usr/lib/libreoffice/program/soffice",
        "/opt/libreoffice/program/soffice",
        "/opt/libreoffice7.6/program/soffice",
        "/opt/libreoffice7.5/program/soffice",
        "/opt/libreoffice7.3/program/soffice",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.name in ("libreoffice", "soffice") and shutil.which(candidate):
            return [candidate, "--headless", "--convert-to"]
        if path.is_file():
            return [str(path), "--headless", "--convert-to"]

    import glob
    for pattern in (
        "/opt/libreoffice*/program/soffice",
        "/opt/apps/*/files/LibreOffice/program/soffice",
    ):
        for match in sorted(glob.glob(pattern)):
            if Path(match).is_file():
                return [match, "--headless", "--convert-to"]
    return None


def _run_libreoffice_convert(source: Path, outdir: Path, target_format: str) -> Path | None:
    converter = _find_libreoffice()
    if converter is None:
        return None
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(
            [*converter, target_format, "--outdir", str(outdir), str(source)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=180,
        )
    except (subprocess.TimeoutExpired, subprocess.CalledProcessError):
        return None

    if target_format.startswith("xls"):
        matches = sorted(outdir.glob(f"{source.stem}*.xls"))
    else:
        matches = sorted(outdir.glob(f"{source.stem}*.xlsx"))
    return matches[0] if matches else None


def _convert_with_libreoffice(source: Path, destination: Path) -> Path:
    converter = _find_libreoffice()
    if converter is None:
        raise ValueError(_format_hint(source, _read_header(source)))

    print(f"正在转换题库格式: {source.name} （WPS → 标准 Excel，约 30 秒）", flush=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        converted = _run_libreoffice_convert(source, tmp_path, "xls:MS Excel 97")
        if converted is None:
            converted = _run_libreoffice_convert(source, tmp_path, "xlsx")
            if converted is not None:
                sheet = _load_openpyxl(converted)
                _write_standard_xls(sheet._rows, destination)
                if is_standard_xls(destination):
                    return destination

        if converted is None:
            if _find_libreoffice() is None:
                raise ValueError(_format_hint(source, _read_header(source)))
            raise ValueError(
                f"无法转换 {source.name}（LibreOffice 无法读取 WPS 格式）。\n"
                "请尝试：\n"
                "  1. 运行: ./转换题库格式.sh\n"
                "  2. 用 LibreOffice 表格打开 → 另存为 → Excel 97-2003 (.xls)\n"
                "  3. 若未安装 LibreOffice，请从麒麟软件中心安装"
            )

        shutil.copy2(converted, destination)

    if not is_standard_xls(destination):
        sheet = _try_read_any(source)
        if sheet is not None:
            _write_standard_xls(sheet._rows, destination)
            return destination
        raise ValueError(f"转换后 {source.name} 仍不是标准 Excel 格式。")
    return destination


def _try_read_any(path: Path) -> SheetAdapter | None:
    if is_standard_xls(path):
        try:
            sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
            rows = [
                [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)
            ]
            return _rows_to_sheet(rows)
        except XLRDError:
            pass

    sheet = _load_calamine(path)
    if sheet is not None:
        return sheet

    sheet = _try_openpyxl_variants(path)
    if sheet is not None:
        return sheet

    header = _read_header(path)
    if header.startswith(b"PK"):
        return None

    text_start = path.read_text(encoding="utf-8", errors="replace")[:200].lower()
    if text_start.lstrip().startswith("<?xml") or "schemas-microsoft-com:office:spreadsheet" in text_start:
        return _load_spreadsheet_xml(path)
    if "<html" in text_start or "<table" in text_start:
        return _load_html_table(path)
    return None


def ensure_readable_excel(path: str | Path) -> Path:
    """确保返回可读路径；WPS/麒麟格式自动读取或转换并缓存。"""
    file_path = Path(path).resolve()
    if not file_path.exists():
        raise FileNotFoundError(str(file_path))

    if is_standard_xls(file_path):
        return file_path

    cached = _cache_path(file_path)
    if cached.exists() and is_standard_xls(cached):
        return cached

    if is_zip_spreadsheet(file_path):
        sheet = _try_openpyxl_variants(file_path)
        if sheet is not None:
            _write_standard_xls(sheet._rows, cached)
            return cached

    sheet = _try_read_any(file_path)
    if sheet is not None:
        _write_standard_xls(sheet._rows, cached)
        return cached

    return _convert_with_libreoffice(file_path, cached)


def _format_hint(path: Path, header: bytes) -> str:
    return "\n".join(
        [
            f"题库文件 {path} 是 WPS/麒麟格式（文件头 {header[:4].hex()}），程序无法直接读取。",
            "",
            "请任选一种方式：",
            "  1. 运行: ./转换题库格式.sh",
            "  2. 安装 LibreOffice 后重新 ./组卷.sh（程序会自动转换）",
            "  3. LibreOffice 打开 → 另存为 → Microsoft Excel 97-2003 (.xls)",
            "",
            "四个文件: 单选.xls 多选.xls 判断.xls 问答.xls",
        ]
    )


def _load_xlrd(path: Path):
    readable = ensure_readable_excel(path)
    try:
        return xlrd.open_workbook(str(readable)).sheet_by_index(0)
    except XLRDError as exc:
        raise ValueError(
            f"{path} 无法读取。\n原始错误: {exc}\n\n请运行 ./转换题库格式.sh"
        ) from exc


def _load_spreadsheet_xml(path: Path) -> SheetAdapter:
    text = path.read_text(encoding="utf-8", errors="replace")
    root = ET.fromstring(text)
    ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
    rows: list[list[Any]] = []
    table = root.find(".//ss:Worksheet/ss:Table", ns) or root.find(
        ".//{urn:schemas-microsoft-com:office:spreadsheet}Table"
    )
    if table is None:
        raise ValueError(_format_hint(path, text[:8].encode("utf-8", errors="replace")))
    for row_el in table.findall("ss:Row", ns) or table.findall(
        "{urn:schemas-microsoft-com:office:spreadsheet}Row"
    ):
        row: list[Any] = []
        for cell in row_el.findall("ss:Cell", ns) or row_el.findall(
            "{urn:schemas-microsoft-com:office:spreadsheet}Cell"
        ):
            data = cell.find("ss:Data", ns) or cell.find(
                "{urn:schemas-microsoft-com:office:spreadsheet}Data"
            )
            row.append("" if data is None or data.text is None else data.text)
        rows.append(row)
    return _rows_to_sheet(rows)


class _TableParser(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._in_table = False
        self._in_row = False
        self._in_cell = False
        self._current_table: list[list[str]] = []
        self._current_row: list[str] = []
        self._cell_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag == "table":
            self._in_table = True
            self._current_table = []
        elif self._in_table and tag == "tr":
            self._in_row = True
            self._current_row = []
        elif self._in_row and tag in ("td", "th"):
            self._in_cell = True
            self._cell_parts = []

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in ("td", "th") and self._in_cell:
            self._current_row.append("".join(self._cell_parts).strip())
            self._in_cell = False
        elif tag == "tr" and self._in_row:
            if self._current_row:
                self._current_table.append(self._current_row)
            self._in_row = False
        elif tag == "table" and self._in_table:
            if self._current_table:
                self.tables.append(self._current_table)
            self._in_table = False

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._cell_parts.append(data)


def _load_html_table(path: Path) -> SheetAdapter:
    text = path.read_text(encoding="utf-8", errors="replace")
    parser = _TableParser()
    parser.feed(text)
    if not parser.tables:
        raise ValueError(_format_hint(path, text[:8].encode("utf-8", errors="replace")))
    return _rows_to_sheet(parser.tables[0])


def open_sheet(path: str):
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(path)

    if is_standard_xls(file_path):
        return _load_xlrd(file_path)

    if is_zip_spreadsheet(file_path):
        sheet = _try_openpyxl_variants(file_path)
        if sheet is not None:
            return sheet
        if zipfile.is_zipfile(file_path):
            return _load_xlrd(file_path)

    text_start = file_path.read_text(encoding="utf-8", errors="replace")[:200].lower()
    if text_start.lstrip().startswith("<?xml") or "schemas-microsoft-com:office:spreadsheet" in text_start:
        return _load_spreadsheet_xml(file_path)
    if "<html" in text_start or "<table" in text_start:
        return _load_html_table(file_path)

    sheet = _try_read_any(file_path)
    if sheet is not None:
        return sheet

    return _load_xlrd(file_path)


def cell_str(sheet, row: int, col: int) -> str:
    value = sheet.cell_value(row, col)
    if value is None:
        return ""
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value).strip()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace").strip()
    return str(value).strip()


def require_columns(headers: list[str], required: list[str], path: str) -> None:
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"{path} 缺少列: {', '.join(missing)}")


def find_column(headers: list[str], name: str) -> int | None:
    target = name.strip().replace(" ", "")
    for idx, header in enumerate(headers):
        if header.strip().replace(" ", "") == target:
            return idx
    return None


def inspect_bank_file(path: str | Path) -> str:
    file_path = Path(path)
    if not file_path.exists():
        return f"缺少  {file_path.name}"
    if is_standard_xls(file_path):
        return f"OK    {file_path.name}  (标准 Excel 97-2003 .xls)"
    if is_zip_spreadsheet(file_path):
        return f"可读  {file_path.name}  (WPS/xlsx 格式，程序可自动读取)"
    if is_wps_binary(file_path):
        lo = "有" if _find_libreoffice() else "无"
        return f"WPS   {file_path.name}  (麒麟 WPS 格式，组卷时自动转换，LibreOffice {lo})"
    return f"待转换 {file_path.name}  (文件头 {_read_header(file_path, 4).hex()})"
