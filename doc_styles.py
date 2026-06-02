"""试卷字体：大题说明宋体五号加粗；题干/选项/答案宋体五号不加粗。"""

from __future__ import annotations

import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

FONT_NAME = "宋体"
SIZE_WUHAO_HALF = 21  # 五号 = 10.5pt


def set_run_font(
    run,
    size_half: int = SIZE_WUHAO_HALF,
    bold: bool = False,
) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_half / 2)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    for tag in ("sz", "szCs"):
        for old in r_pr.findall(qn(f"w:{tag}")):
            r_pr.remove(old)
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:val"), str(size_half))
        r_pr.append(node)
    for old in r_pr.findall(qn("w:b")):
        r_pr.remove(old)
    if bold:
        r_pr.append(OxmlElement("w:b"))


def add_text_run(paragraph: Paragraph, text: str, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_stem_paragraph(subdoc, text: str) -> Paragraph:
    paragraph = subdoc.add_paragraph()
    add_text_run(paragraph, text, bold=False)
    return paragraph


def add_options_paragraph(subdoc, text: str) -> Paragraph:
    paragraph = subdoc.add_paragraph()
    add_text_run(paragraph, text, bold=False)
    return paragraph


def add_answer_line_paragraph(subdoc, text: str) -> Paragraph:
    paragraph = subdoc.add_paragraph()
    add_text_run(paragraph, text, bold=False)
    return paragraph


def add_body_paragraph(subdoc, text: str) -> Paragraph:
    paragraph = subdoc.add_paragraph()
    add_text_run(paragraph, text, bold=False)
    return paragraph


def add_blank_paragraph(subdoc) -> Paragraph:
    return subdoc.add_paragraph("")
